"""Xuất báo cáo Word (.docx) tự động từ một run trong registry."""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

from src.models.registry import RunArtifacts  # noqa: E402
from src.models.segmentation import SEGMENT_INFO, base_segment  # noqa: E402


def _fig_to_stream(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _table(doc: Document, df: pd.DataFrame, fmt: dict | None = None) -> None:
    fmt = fmt or {}
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for i, col in enumerate(df.columns):
        t.rows[0].cells[i].text = str(col)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, col in enumerate(df.columns):
            v = row[col]
            cells[i].text = fmt[col].format(v) if col in fmt and pd.notna(v) else str(v)


def build_report(art: RunArtifacts, out_dir: Path) -> Path:
    md = art.metadata
    c = art.customers
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading("BÁO CÁO PHÂN TÍCH & PHÂN KHÚC KHÁCH HÀNG", 0)
    doc.add_paragraph(f"Mã lần chạy: {art.run_id}  |  Ngày tạo: {datetime.now():%d/%m/%Y %H:%M}  |  "
                      f"Ngày mốc dữ liệu: {md['snapshot_date']}")

    doc.add_heading("1. Tổng quan dữ liệu", 1)
    d = md["data"]
    _table(doc, pd.DataFrame([
        ["Số dòng giao dịch", f"{d['rows']:,}"], ["Số khách hàng", f"{d['customers']:,}"],
        ["Số mã giao dịch", f"{d['transactions']:,}"], ["Khoảng thời gian", f"{d['date_min']} → {d['date_max']}"],
        ["% mã giao dịch dùng chung nhiều khách", f"{d['shared_transaction_pct']:.1f}%"],
        ["Cách tính Frequency", "Số ngày mua khác nhau" if md["frequency_by_days"] else "Số mã giao dịch"],
    ], columns=["Chỉ tiêu", "Giá trị"]))

    doc.add_heading("2. Lựa chọn số cụm", 1)
    ks = pd.DataFrame(md["k_search"])
    fig, ax = plt.subplots(1, 2, figsize=(10, 3.2))
    ax[0].plot(ks["k"], ks["wcss"], marker="o"); ax[0].set_title("Elbow (WCSS)"); ax[0].set_xlabel("k")
    ax[1].plot(ks["k"], ks["silhouette"], marker="o", color="orange"); ax[1].set_title("Silhouette")
    ax[1].set_xlabel("k")
    doc.add_picture(_fig_to_stream(fig), width=Inches(6.3))
    m = md["segmentation_metrics"]
    doc.add_paragraph(f"Số cụm được chọn: k = {md['k']} (Silhouette tốt nhất ở k = {md['best_k_silhouette']}). "
                      f"Silhouette = {m['silhouette']:.3f}, Davies-Bouldin = {m['davies_bouldin']:.3f}, "
                      f"Calinski-Harabasz = {m['calinski_harabasz']:,.1f}.")

    if md.get("algorithm_comparison"):
        doc.add_heading("3. So sánh thuật toán phân cụm (cùng một mẫu đánh giá)", 1)
        comp = pd.DataFrame(md["algorithm_comparison"])
        cols = [x for x in ["algorithm", "n_clusters", "noise_pct", "silhouette", "davies_bouldin",
                            "calinski_harabasz", "dunn"] if x in comp.columns]
        _table(doc, comp[cols], {"noise_pct": "{:.1f}", "silhouette": "{:.3f}", "davies_bouldin": "{:.3f}",
                                 "calinski_harabasz": "{:,.1f}", "dunn": "{:.3f}"})

    doc.add_heading("4. Đặc điểm các phân khúc", 1)
    prof = pd.DataFrame(md["segment_profile"])
    _table(doc, prof[["Segment", "customers", "customer_pct", "recency_mean", "frequency_mean",
                      "monetary_mean", "revenue_pct"]].rename(columns={
        "Segment": "Phân khúc", "customers": "Số KH", "customer_pct": "% KH", "recency_mean": "R TB",
        "frequency_mean": "F TB", "monetary_mean": "M TB", "revenue_pct": "% Doanh thu"}),
        {"Số KH": "{:,}", "% KH": "{:.1f}", "R TB": "{:.0f}", "F TB": "{:.2f}", "M TB": "{:,.0f}",
         "% Doanh thu": "{:.1f}"})
    fig, ax = plt.subplots(figsize=(7, 4))
    for seg, g in c.sample(min(8000, len(c)), random_state=1).groupby("Segment"):
        ax.scatter(g["Recency"], g["Monetary"], s=6, alpha=0.5, label=seg)
    ax.set_xlabel("Recency (ngày)"); ax.set_ylabel("Monetary"); ax.legend(markerscale=3)
    doc.add_picture(_fig_to_stream(fig), width=Inches(5.8))
    for seg in prof["Segment"]:
        doc.add_paragraph(f"{seg}: {SEGMENT_INFO.get(base_segment(seg), '')}", style="List Bullet")

    doc.add_heading("5. Mô hình dự báo", 1)
    pm = md.get("predictor_metrics")
    if pm:
        _table(doc, pd.DataFrame([
            ["Mốc cắt huấn luyện", pm["cutoff"]], ["Cửa sổ dự báo", f"{pm['horizon_days']} ngày"],
            ["Tỷ lệ churn thực tế", f"{pm['churn_rate']:.1%}"], ["ROC-AUC (churn)", f"{pm['churn_auc']:.3f}"],
            ["Accuracy / mức cơ sở", f"{pm['churn_accuracy']:.3f} / {pm['churn_baseline_accuracy']:.3f}"],
            ["MAE giá trị / mức cơ sở", f"{pm['value_mae']:,.2f} / {pm['value_baseline_mae']:,.2f}"],
        ], columns=["Chỉ tiêu", "Giá trị"]))
        if not pm.get("churn_reliable"):
            doc.add_paragraph("Kết luận: mô hình dự báo KHÔNG vượt mức ngẫu nhiên (AUC < 0,60) nên không được "
                              "dùng cho ra quyết định. Nguyên nhân khả dĩ: dữ liệu không chứa tín hiệu hành vi "
                              "lặp lại (vd dữ liệu tổng hợp/sinh ngẫu nhiên).")
    else:
        doc.add_paragraph(f"Không huấn luyện được: {md.get('predictor_error')}")

    doc.add_heading("6. Khuyến nghị hành động", 1)
    from src.ai_agent.copilot import PLAYBOOK
    for seg in prof["Segment"]:
        play = PLAYBOOK.get(base_segment(seg))
        if play:
            doc.add_paragraph(f"{seg}: {play['insight']} Ưu đãi gợi ý: {play['offer'].split('{')[0].strip()}",
                              style="List Number")

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"Bao_cao_{art.run_id}.docx"
    doc.save(path)
    return path
